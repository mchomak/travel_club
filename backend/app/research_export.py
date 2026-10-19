from __future__ import annotations

from collections.abc import Iterable
from datetime import date, datetime
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.models import ResearchAnswer, ResearchInterview, ResearchRespondent
from app.research_bank import load_question_bank, load_screening

EXCEL_CELL_LIMIT = 32_767


def _iso(value: datetime | date | None) -> str | None:
    return value.isoformat() if value is not None else None


def _display(value: Any) -> str | int | bool | None:
    if value is None or isinstance(value, (str, int, bool)):
        return value
    return str(value)


def _append_row(sheet, values: Iterable[Any]) -> None:
    row = 1 if sheet.max_row == 1 and sheet.cell(1, 1).value is None else sheet.max_row + 1
    for column, value in enumerate(values, start=1):
        cell = sheet.cell(row=row, column=column, value=_display(value))
        if isinstance(value, str):
            # openpyxl infers formulas from leading '='. Force interview text to string data.
            cell.data_type = "s"
        cell.alignment = Alignment(vertical="top", wrap_text=True)


def _snapshot_questions(interview: ResearchInterview) -> list[dict[str, Any]]:
    if isinstance(interview.question_snapshot, list):
        return [item for item in interview.question_snapshot if isinstance(item, dict)]
    return []


def build_research_workbook(
    respondents: list[ResearchRespondent],
    interviews: list[ResearchInterview],
    answers: list[ResearchAnswer],
) -> bytes:
    """Build a complete, formula-safe Excel export from an already consistent DB snapshot."""
    workbook = Workbook()
    respondents_sheet = workbook.active
    respondents_sheet.title = "Respondents"
    interviews_sheet = workbook.create_sheet("Interviews")
    answers_sheet = workbook.create_sheet("Answers")
    questions_sheet = workbook.create_sheet("Questions")
    hypotheses_sheet = workbook.create_sheet("Hypotheses")

    screening_fields = load_screening().get("fields", [])
    screening_keys = [
        (field["key"], field.get("label", field["key"]))
        for field in screening_fields
        if isinstance(field, dict) and isinstance(field.get("key"), str)
    ]
    known_screening = {key for key, _ in screening_keys}
    historical_screening = sorted(
        {
            key
            for respondent in respondents
            for key in (respondent.screening_answers or {})
            if isinstance(key, str) and key not in known_screening
        }
    )
    screening_keys.extend((key, key) for key in historical_screening)
    respondent_headers = [
        "Respondent ID",
        "Name",
        "Age / age range",
        "City",
        "Suggested category",
        "Final category",
        "Category override",
        "Research goal",
        "Status",
        "Created At",
        "Updated At",
        *[f"Screening: {label}" for _, label in screening_keys],
    ]
    _append_row(respondents_sheet, respondent_headers)

    latest_interview: dict[Any, ResearchInterview] = {}
    for interview in interviews:
        current = latest_interview.get(interview.respondent_id)
        if current is None or interview.created_at > current.created_at:
            latest_interview[interview.respondent_id] = interview
    for respondent in respondents:
        interview = latest_interview.get(respondent.id)
        screening = respondent.screening_answers or {}
        _append_row(
            respondents_sheet,
            [
                str(respondent.id),
                respondent.name,
                respondent.age_range,
                respondent.city,
                respondent.suggested_category,
                respondent.final_category,
                respondent.category_override,
                respondent.research_goal,
                interview.status if interview else None,
                _iso(respondent.created_at),
                _iso(respondent.updated_at),
                *[screening.get(key) for key, _ in screening_keys],
            ],
        )

    _append_row(
        interviews_sheet,
        [
            "Interview ID",
            "Respondent ID",
            "Name",
            "Final category",
            "Scenario code",
            "Scenario version",
            "Bank version",
            "Status",
            "Started At",
            "Completed At",
            "Created At",
            "Updated At",
        ],
    )
    respondents_by_id = {respondent.id: respondent for respondent in respondents}
    for interview in interviews:
        respondent = respondents_by_id.get(interview.respondent_id)
        _append_row(
            interviews_sheet,
            [
                str(interview.id),
                str(interview.respondent_id),
                respondent.name if respondent else None,
                respondent.final_category if respondent else None,
                interview.scenario_code,
                interview.scenario_version,
                interview.bank_version,
                interview.status,
                _iso(interview.started_at),
                _iso(interview.completed_at),
                _iso(interview.created_at),
                _iso(interview.updated_at),
            ],
        )

    answers_by_interview: dict[Any, dict[str, ResearchAnswer]] = {}
    for answer in answers:
        answers_by_interview.setdefault(answer.interview_id, {})[answer.question_code] = answer

    answer_records: list[
        tuple[ResearchInterview, ResearchRespondent | None, dict[str, Any], ResearchAnswer | None]
    ] = []
    question_versions: dict[tuple[str, str], dict[str, Any]] = {}
    max_answer_parts = 1
    for interview in interviews:
        respondent = respondents_by_id.get(interview.respondent_id)
        saved = answers_by_interview.get(interview.id, {})
        for question in _snapshot_questions(interview):
            code = question.get("code")
            if not isinstance(code, str):
                continue
            answer = saved.get(code)
            answer_records.append((interview, respondent, question, answer))
            question_bank_version = str(question.get("bank_version") or interview.bank_version)
            key = (question_bank_version, code)
            question_versions.setdefault(key, {**question, "bank_version": question_bank_version})
            text = answer.answer_text if answer else ""
            max_answer_parts = max(
                max_answer_parts, (len(text) + EXCEL_CELL_LIMIT - 1) // EXCEL_CELL_LIMIT
            )

    _append_row(
        answers_sheet,
        [
            "Respondent ID",
            "Name",
            "Category",
            "Interview ID",
            "Question Code",
            "Question Text",
            "Answer Text",
            *[f"Answer Text (part {part})" for part in range(2, max_answer_parts + 1)],
            "Hypothesis IDs",
            "Updated At",
        ],
    )
    for interview, respondent, question, answer in answer_records:
        answer_text = answer.answer_text if answer else ""
        chunks = [
            answer_text[index : index + EXCEL_CELL_LIMIT]
            for index in range(0, len(answer_text), EXCEL_CELL_LIMIT)
        ]
        _append_row(
            answers_sheet,
            [
                str(respondent.id) if respondent else str(interview.respondent_id),
                respondent.name if respondent else None,
                respondent.final_category if respondent else None,
                str(interview.id),
                question.get("code"),
                question.get("text"),
                *(chunks + [""] * max(0, max_answer_parts - len(chunks))),
                ", ".join(question.get("hypothesis_ids", [])),
                _iso(answer.updated_at) if answer else None,
            ],
        )

    _append_row(
        questions_sheet,
        ["Bank Version", "Question Code", "Section", "Question Text", "Type", "Hypothesis IDs"],
    )
    for key in sorted(question_versions):
        question = question_versions[key]
        _append_row(
            questions_sheet,
            [
                question["bank_version"],
                question.get("code"),
                question.get("section"),
                question.get("text"),
                question.get("type"),
                ", ".join(question.get("hypothesis_ids", [])),
            ],
        )

    _append_row(hypotheses_sheet, ["Hypothesis ID", "Description"])
    for hypothesis in load_question_bank().get("hypotheses", []):
        if isinstance(hypothesis, dict):
            _append_row(hypotheses_sheet, [hypothesis.get("code"), hypothesis.get("text")])

    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="173F7A")
        for column_cells in sheet.columns:
            width = min(max(len(str(cell.value or "")) for cell in column_cells) + 2, 48)
            sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = width

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
