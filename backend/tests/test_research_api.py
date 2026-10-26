from __future__ import annotations

from concurrent.futures import ThreadPoolExecutor
from copy import deepcopy
from datetime import UTC, datetime, timedelta
from io import BytesIO
from uuid import UUID, uuid4

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import delete, event

from app import research as research_api
from app.auth import token_digest
from app.config import get_settings
from app.db import SessionLocal, get_db
from app.main import create_app
from app.models import (
    ResearchInterview,
    ResearchInterviewVersion,
    ResearchRespondent,
    User,
)
from app.models import Session as AdminSession
from tests.conftest import FIXTURE_PASSWORD

HEADERS = {"X-Requested-With": "TravelClub", "Origin": "http://localhost:3000"}


@pytest.fixture
def authenticated(client, admin):
    response = client.post(
        "/api/auth/login",
        headers=HEADERS,
        json={"login": admin.login, "password": FIXTURE_PASSWORD},
    )
    assert response.status_code == 200
    return client


def test_research_respondent_api_requires_admin_session(client):
    for path in (
        "/api/research/respondents",
        "/api/research/config",
        "/api/research/questions",
        "/api/research/export.xlsx",
        "/api/research/interviews/00000000-0000-0000-0000-000000000001/versions",
        "/api/research/interviews/00000000-0000-0000-0000-000000000001/versions/"
        "00000000-0000-0000-0000-000000000002",
    ):
        response = client.get(path, headers=HEADERS)
        assert response.status_code == 401, path
        assert response.json()["message"]


def test_research_origin_is_allowed_only_when_configured(authenticated, monkeypatch):
    settings = get_settings()
    monkeypatch.setattr(settings, "research_origin", "https://research.example.test")
    body = {"screening_answers": {"name": "Configured origin"}}

    accepted = authenticated.post(
        "/api/research/respondents",
        headers={**HEADERS, "Origin": "https://research.example.test"},
        json=body,
    )
    assert accepted.status_code == 201

    rejected = authenticated.post(
        "/api/research/respondents",
        headers={**HEADERS, "Origin": "https://other.example.test"},
        json=body,
    )
    assert rejected.status_code == 403
    assert rejected.json()["code"] == "CSRF_REJECTED"


def test_authenticated_configuration_endpoint_exposes_editable_questionnaire(authenticated):
    response = authenticated.get("/api/research/config")

    assert response.status_code == 200
    configuration = response.json()
    assert configuration["validation_issues"] == []
    assert len(configuration["bank"]["questions"]) == 140
    assert len(configuration["bank"]["hypotheses"]) == 28
    assert {module["code"] for module in configuration["scenarios"]["expanded_modules"]} == {
        "A", "B", "C", "D", "E", "G", "H", "I", "J", "L"
    }
    assert len(configuration["screening"]["fields"]) == 9
    assert configuration["offers_enabled"] is False


def test_config_reports_malformed_sections_without_dereferencing_them(authenticated, monkeypatch):
    monkeypatch.setattr(
        research_api,
        "load_configuration",
        lambda: {"bank": {}, "scenarios": None, "screening": None, "classification": None},
    )

    response = authenticated.get("/api/research/config")

    assert response.status_code == 200
    assert "scenarios JSON root must be an object" in response.json()["validation_issues"]
    assert response.json()["offers_enabled"] is False


def test_old_interview_keeps_question_text_after_json_bank_changes(authenticated, monkeypatch):
    original_configuration = research_api.load_configuration()
    first = authenticated.post(
        "/api/research/respondents",
        headers=HEADERS,
        json={
            "screening_answers": {"name": "Банк версия один", "independent_trips_24m": "3_plus"}
        },
    ).json()
    first_id = first["interviews"][0]["id"]
    started = authenticated.post(
        f"/api/research/interviews/{first_id}/start", headers=HEADERS
    ).json()
    original_text = next(q["text"] for q in started["question_snapshot"] if q["code"] == "G01")

    changed = deepcopy(original_configuration)
    next(q for q in changed["bank"]["questions"] if q["code"] == "G01")[
        "text"
    ] = "Новая формулировка"
    monkeypatch.setattr(research_api, "load_configuration", lambda: changed)

    old_detail = authenticated.get(f"/api/research/respondents/{first['id']}").json()
    old_question = next(
        q for q in old_detail["interviews"][0]["question_snapshot"] if q["code"] == "G01"
    )
    assert old_question["text"] == original_text

    second = authenticated.post(
        "/api/research/respondents",
        headers=HEADERS,
        json={
            "screening_answers": {"name": "Банк версия два", "independent_trips_24m": "3_plus"}
        },
    ).json()
    second_started = authenticated.post(
        f"/api/research/interviews/{second['interviews'][0]['id']}/start", headers=HEADERS
    ).json()
    new_question = next(
        q for q in second_started["question_snapshot"] if q["code"] == "G01"
    )
    assert new_question["text"] == "Новая формулировка"


def test_create_respondent_saves_screening_and_draft_interview_atomically(authenticated):
    response = authenticated.post(
        "/api/research/respondents",
        headers=HEADERS,
        json={"screening_answers": {"name": "Мария", "age_range": "26–35", "city": "Казань"}},
    )

    assert response.status_code == 201, response.text
    created = response.json()
    assert created["screening_answers"]["name"] == "Мария"
    assert created["age_range"] == "26–35"
    assert created["city"] == "Казань"
    assert created["suggested_category"] == created["final_category"]
    assert created["category_override"] is False
    assert len(created["interviews"]) == 1
    assert created["interviews"][0]["status"] == "draft"
    assert created["interviews"][0]["question_snapshot"] == []


def test_manual_category_override_is_snapshotted_when_interview_starts(authenticated):
    created = authenticated.post(
        "/api/research/respondents",
        headers=HEADERS,
        json={"screening_answers": {"name": "Василиса"}},
    ).json()

    selected = authenticated.patch(
        f"/api/research/respondents/{created['id']}/category",
        headers=HEADERS,
        json={"final_category": "travel_ready"},
    )
    assert selected.status_code == 200
    assert selected.json()["suggested_category"] != "travel_ready"
    assert selected.json()["final_category"] == "travel_ready"
    assert selected.json()["category_override"] is True

    interview_id = created["interviews"][0]["id"]
    started = authenticated.post(f"/api/research/interviews/{interview_id}/start", headers=HEADERS)
    assert started.status_code == 200
    interview = started.json()
    assert interview["status"] == "in_progress"
    assert interview["scenario_code"] == "travel_ready"
    assert [question["code"] for question in interview["question_snapshot"]] == [
        "G01",
        "G04",
        "G07",
        "G10",
        "R01",
        "R02",
        "R03",
        "R05",
        "R06",
        "R12",
    ]
    assert all(
        question["text"] and question["bank_version"] for question in interview["question_snapshot"]
    )


def test_answer_compare_and_swap_is_idempotent_and_preserves_newer_text(authenticated):
    created = authenticated.post(
        "/api/research/respondents",
        headers=HEADERS,
        json={"screening_answers": {"name": "Ирина"}},
    ).json()
    interview_id = created["interviews"][0]["id"]
    authenticated.post(f"/api/research/interviews/{interview_id}/start", headers=HEADERS)
    path = f"/api/research/interviews/{interview_id}/answers/G01"

    first = authenticated.put(
        path, headers=HEADERS, json={"answer_text": "Первая заметка", "expected_revision": 0}
    )
    assert first.status_code == 200
    assert first.json()["revision"] == 1
    retry = authenticated.put(
        path, headers=HEADERS, json={"answer_text": "Первая заметка", "expected_revision": 0}
    )
    assert retry.status_code == 200
    assert retry.json()["revision"] == 1

    stale = authenticated.put(
        path, headers=HEADERS, json={"answer_text": "Устаревший текст", "expected_revision": 0}
    )
    assert stale.status_code == 409
    assert stale.json()["current"]["answer_text"] == "Первая заметка"
    assert stale.json()["current"]["revision"] == 1

    latest = authenticated.put(
        path, headers=HEADERS, json={"answer_text": "Обновлённая заметка", "expected_revision": 1}
    )
    assert latest.status_code == 200
    assert latest.json()["answer_text"] == "Обновлённая заметка"
    assert latest.json()["revision"] == 2


def test_completion_persists_initial_readable_version(authenticated):
    created = authenticated.post(
        "/api/research/respondents",
        headers=HEADERS,
        json={"screening_answers": {"name": "Versioned respondent", "city": "Lisbon"}},
    ).json()
    interview_id = created["interviews"][0]["id"]
    started = authenticated.post(
        f"/api/research/interviews/{interview_id}/start", headers=HEADERS
    )
    assert started.status_code == 200
    saved = authenticated.put(
        f"/api/research/interviews/{interview_id}/answers/G01",
        headers=HEADERS,
        json={"answer_text": "The original answer", "expected_revision": 0},
    )
    assert saved.status_code == 200
    completed = authenticated.post(
        f"/api/research/interviews/{interview_id}/complete", headers=HEADERS
    )
    assert completed.status_code == 200

    listed = authenticated.get(
        f"/api/research/interviews/{interview_id}/versions", headers=HEADERS
    )

    assert listed.status_code == 200
    versions = listed.json()
    assert versions["latest_version_number"] == 1
    assert [item["version_number"] for item in versions["items"]] == [1]
    first = versions["items"][0]
    assert first["created_at"].endswith("Z")

    detail = authenticated.get(
        f"/api/research/interviews/{interview_id}/versions/{first['id']}", headers=HEADERS
    )

    assert detail.status_code == 200
    snapshot = detail.json()["snapshot"]
    assert snapshot["respondent"]["name"] == "Versioned respondent"
    assert snapshot["respondent"]["city"] == "Lisbon"
    assert snapshot["interview"]["status"] == "completed"
    assert any(
        question["code"] == "G01" and question["text"]
        for question in snapshot["interview"]["question_snapshot"]
    )
    assert snapshot["interview"]["answers"]["G01"]["answer_text"] == "The original answer"
    assert snapshot["interview"]["answers"]["G01"]["revision"] == 1


def test_manual_version_save_updates_answers_and_preserves_history(authenticated):
    created = authenticated.post(
        "/api/research/respondents",
        headers=HEADERS,
        json={
            "screening_answers": {
                "name": "Version edit",
                "independent_trips_24m": "3_plus",
            }
        },
    ).json()
    interview_id = created["interviews"][0]["id"]
    started = authenticated.post(
        f"/api/research/interviews/{interview_id}/start", headers=HEADERS
    ).json()
    original = authenticated.put(
        f"/api/research/interviews/{interview_id}/answers/G01",
        headers=HEADERS,
        json={"answer_text": "Original answer", "expected_revision": 0},
    )
    assert original.status_code == 200
    authenticated.post(f"/api/research/interviews/{interview_id}/complete", headers=HEADERS)
    first_id = authenticated.get(
        f"/api/research/interviews/{interview_id}/versions", headers=HEADERS
    ).json()["items"][0]["id"]

    answers = {question["code"]: "" for question in started["question_snapshot"]}
    answers["G01"] = "Revised answer"
    answers["G02"] = "Follow-up answer"
    saved = authenticated.post(
        f"/api/research/interviews/{interview_id}/versions",
        headers=HEADERS,
        json={"answers": answers},
    )

    assert saved.status_code == 201
    second = saved.json()
    assert second["version_number"] == 2
    assert second["snapshot"]["interview"]["answers"]["G01"]["answer_text"] == "Revised answer"
    assert second["snapshot"]["interview"]["answers"]["G01"]["revision"] == 2
    assert second["snapshot"]["interview"]["answers"]["G02"]["revision"] == 1
    assert second["created_at"].endswith("Z")

    first = authenticated.get(
        f"/api/research/interviews/{interview_id}/versions/{first_id}", headers=HEADERS
    ).json()
    assert first["snapshot"]["interview"]["answers"]["G01"]["answer_text"] == "Original answer"
    assert first["snapshot"]["interview"]["answers"]["G01"]["revision"] == 1

    repeated = authenticated.post(
        f"/api/research/interviews/{interview_id}/versions",
        headers=HEADERS,
        json={"answers": answers},
    )
    assert repeated.status_code == 201
    assert repeated.json()["version_number"] == 3
    assert repeated.json()["snapshot"]["interview"]["answers"]["G01"]["revision"] == 2
    history = authenticated.get(
        f"/api/research/interviews/{interview_id}/versions", headers=HEADERS
    ).json()
    assert [item["version_number"] for item in history["items"]] == [3, 2, 1]
    assert history["latest_version_number"] == 3


def test_manual_version_save_rolls_back_answers_when_snapshot_commit_fails(authenticated, db):
    created = authenticated.post(
        "/api/research/respondents",
        headers=HEADERS,
        json={"screening_answers": {"name": "Rollback version"}},
    ).json()
    interview_id = created["interviews"][0]["id"]
    started = authenticated.post(
        f"/api/research/interviews/{interview_id}/start", headers=HEADERS
    ).json()
    authenticated.put(
        f"/api/research/interviews/{interview_id}/answers/G01",
        headers=HEADERS,
        json={"answer_text": "Persisted answer", "expected_revision": 0},
    )
    authenticated.post(f"/api/research/interviews/{interview_id}/complete", headers=HEADERS)
    answers = {question["code"]: "" for question in started["question_snapshot"]}
    answers["G01"] = "Uncommitted replacement"

    def add_conflicting_version(session):
        session.add(
            ResearchInterviewVersion(
                interview_id=UUID(interview_id),
                version_number=2,
                snapshot={},
                created_at=datetime.now(UTC),
            )
        )

    event.listen(db, "before_commit", add_conflicting_version)
    try:
        failed = authenticated.post(
            f"/api/research/interviews/{interview_id}/versions",
            headers=HEADERS,
            json={"answers": answers},
        )
    finally:
        event.remove(db, "before_commit", add_conflicting_version)

    assert failed.status_code == 500
    detail = authenticated.get(f"/api/research/respondents/{created['id']}").json()
    assert detail["interviews"][0]["answers"]["G01"]["answer_text"] == "Persisted answer"
    history = authenticated.get(
        f"/api/research/interviews/{interview_id}/versions", headers=HEADERS
    ).json()
    assert history["latest_version_number"] == 1
    assert [item["version_number"] for item in history["items"]] == [1]


def test_version_endpoints_reject_incomplete_and_unknown_question_codes(authenticated):
    created = authenticated.post(
        "/api/research/respondents",
        headers=HEADERS,
        json={"screening_answers": {"name": "Version validation"}},
    ).json()
    interview_id = created["interviews"][0]["id"]
    versions_path = f"/api/research/interviews/{interview_id}/versions"

    incomplete_list = authenticated.get(versions_path, headers=HEADERS)
    incomplete_save = authenticated.post(
        versions_path, headers=HEADERS, json={"answers": {}}
    )
    assert incomplete_list.status_code == 409
    assert incomplete_save.status_code == 409

    started = authenticated.post(
        f"/api/research/interviews/{interview_id}/start", headers=HEADERS
    ).json()
    authenticated.post(f"/api/research/interviews/{interview_id}/complete", headers=HEADERS)
    answers = {question["code"]: "" for question in started["question_snapshot"]}
    omitted_code = next(iter(answers))
    del answers[omitted_code]
    missing_answer = authenticated.post(
        versions_path, headers=HEADERS, json={"answers": answers}
    )
    unknown = authenticated.post(
        versions_path, headers=HEADERS, json={"answers": {"ZZZ": "Invalid answer"}}
    )
    missing_detail = authenticated.get(
        f"{versions_path}/{uuid4()}", headers=HEADERS
    )

    assert unknown.status_code == 422
    assert unknown.json()["code"] == "UNKNOWN_QUESTION"
    assert missing_answer.status_code == 422
    assert missing_answer.json()["code"] == "ANSWERS_REQUIRED"
    assert missing_answer.json()["missing_codes"] == [omitted_code]
    assert missing_detail.status_code == 404
    assert missing_detail.json()["code"] == "INTERVIEW_VERSION_NOT_FOUND"


def test_completed_interview_rejects_question_changes_and_autosave(authenticated):
    created = authenticated.post(
        "/api/research/respondents",
        headers=HEADERS,
        json={"screening_answers": {"name": "Полина"}},
    ).json()
    interview_id = created["interviews"][0]["id"]
    authenticated.post(f"/api/research/interviews/{interview_id}/start", headers=HEADERS)

    completed = authenticated.post(
        f"/api/research/interviews/{interview_id}/complete", headers=HEADERS
    )
    assert completed.status_code == 200
    assert completed.json()["status"] == "completed"
    assert completed.json()["completed_at"]

    appended = authenticated.post(
        f"/api/research/interviews/{interview_id}/questions",
        headers=HEADERS,
        json={"question_codes": ["B01"]},
    )
    assert appended.status_code == 409
    assert appended.json()["code"] == "INTERVIEW_COMPLETED"
    saved = authenticated.put(
        f"/api/research/interviews/{interview_id}/answers/B01",
        headers=HEADERS,
        json={"answer_text": "Поздняя запись", "expected_revision": 0},
    )
    assert saved.status_code == 409
    assert saved.json()["code"] == "INTERVIEW_COMPLETED"
    detail = authenticated.get(f"/api/research/respondents/{created['id']}").json()
    assert detail["interviews"][0]["status"] == "completed"
    assert "B01" not in detail["interviews"][0]["answers"]


def test_brand_questions_stay_at_end_when_supplementing_in_progress_interview(authenticated):
    created = authenticated.post(
        "/api/research/respondents",
        headers=HEADERS,
        json={"screening_answers": {"name": "Brand order", "independent_trips_24m": "0"}},
    ).json()
    interview_id = created["interviews"][0]["id"]
    authenticated.post(f"/api/research/interviews/{interview_id}/start", headers=HEADERS)

    first_addition = authenticated.post(
        f"/api/research/interviews/{interview_id}/questions",
        headers=HEADERS,
        json={"question_codes": ["B01"]},
    )
    assert first_addition.status_code == 200
    later_addition = authenticated.post(
        f"/api/research/interviews/{interview_id}/questions",
        headers=HEADERS,
        json={"question_codes": ["T01"]},
    )

    assert later_addition.status_code == 200
    codes = [question["code"] for question in later_addition.json()["question_snapshot"]]
    assert codes.index("T01") < codes.index("B01")
    assert all(code.startswith("B") for code in codes[codes.index("B01"):])


def test_new_expanded_question_can_be_appended_with_its_module_metadata(authenticated):
    created = authenticated.post(
        "/api/research/respondents",
        headers=HEADERS,
        json={"screening_answers": {"name": "Expanded module", "independent_trips_24m": "0"}},
    ).json()
    interview_id = created["interviews"][0]["id"]
    authenticated.post(f"/api/research/interviews/{interview_id}/start", headers=HEADERS)

    appended = authenticated.post(
        f"/api/research/interviews/{interview_id}/questions",
        headers=HEADERS,
        json={"question_codes": ["N43"]},
    )

    assert appended.status_code == 200
    question = next(
        item for item in appended.json()["question_snapshot"] if item["code"] == "N43"
    )
    assert question["module_code"] == "H"
    assert question["hypothesis_ids"] == ["H26"]
    assert question["bank_version"] == "2.2"


def test_screening_recalculation_does_not_rewrite_started_route_or_answers(authenticated):
    created = authenticated.post(
        "/api/research/respondents",
        headers=HEADERS,
        json={"screening_answers": {"name": "Анна", "independent_trips_24m": "0"}},
    ).json()
    interview_id = created["interviews"][0]["id"]
    started = authenticated.post(
        f"/api/research/interviews/{interview_id}/start", headers=HEADERS
    ).json()
    authenticated.put(
        f"/api/research/interviews/{interview_id}/answers/G01",
        headers=HEADERS,
        json={"answer_text": "Уже сохранено", "expected_revision": 0},
    )

    edited = authenticated.patch(
        f"/api/research/respondents/{created['id']}/screening",
        headers=HEADERS,
        json={
            "screening_answers": {"name": "Анна", "independent_trips_24m": "3_plus"},
            "apply_suggestion": False,
        },
    )
    assert edited.status_code == 200
    assert edited.json()["suggested_category"] == "trips_experienced"
    assert edited.json()["final_category"] == "trips_beginner"
    interview = edited.json()["interviews"][0]
    assert interview["scenario_code"] == started["scenario_code"] == "trips_beginner"
    assert interview["question_snapshot"] == started["question_snapshot"]
    assert interview["answers"]["G01"]["answer_text"] == "Уже сохранено"


def test_screening_and_category_changes_version_every_completed_interview(authenticated, db):
    created = authenticated.post(
        "/api/research/respondents",
        headers=HEADERS,
        json={"screening_answers": {"name": "Original profile", "independent_trips_24m": "0"}},
    ).json()
    interview_id = created["interviews"][0]["id"]
    started = authenticated.post(
        f"/api/research/interviews/{interview_id}/start", headers=HEADERS
    ).json()
    authenticated.post(f"/api/research/interviews/{interview_id}/complete", headers=HEADERS)
    version_path = f"/api/research/interviews/{interview_id}/versions"
    first_version = authenticated.get(version_path, headers=HEADERS).json()["items"][0]
    original_snapshot = authenticated.get(
        f"{version_path}/{first_version['id']}", headers=HEADERS
    ).json()["snapshot"]

    second_interview = ResearchInterview(
        respondent_id=UUID(created["id"]),
        scenario_code=started["scenario_code"],
        scenario_version=started["scenario_version"],
        bank_version=started["bank_version"],
        status="completed",
        question_snapshot=deepcopy(started["question_snapshot"]),
        completed_at=datetime.now(UTC),
    )
    db.add(second_interview)
    db.flush()
    second_snapshot = deepcopy(original_snapshot)
    second_snapshot["interview"]["id"] = str(second_interview.id)
    db.add(
        ResearchInterviewVersion(
            interview_id=second_interview.id,
            version_number=1,
            snapshot=second_snapshot,
            created_at=datetime.now(UTC),
        )
    )
    db.commit()

    screening = authenticated.patch(
        f"/api/research/respondents/{created['id']}/screening",
        headers=HEADERS,
        json={
            "screening_answers": {
                "name": "Updated profile",
                "city": "Tokyo",
                "independent_trips_24m": "3_plus",
            },
            "apply_suggestion": False,
        },
    )
    assert screening.status_code == 200
    interview_ids = (interview_id, str(second_interview.id))
    histories = [
        authenticated.get(
            f"/api/research/interviews/{item_id}/versions", headers=HEADERS
        ).json()
        for item_id in interview_ids
    ]
    assert [history["latest_version_number"] for history in histories] == [2, 2]
    first_history = authenticated.get(
        f"{version_path}/{first_version['id']}", headers=HEADERS
    ).json()
    assert first_history["snapshot"]["respondent"]["name"] == "Original profile"
    for item_id, history in zip(interview_ids, histories, strict=True):
        latest = history["items"][0]
        snapshot = authenticated.get(
            f"/api/research/interviews/{item_id}/versions/{latest['id']}", headers=HEADERS
        ).json()["snapshot"]
        assert snapshot["respondent"]["name"] == "Updated profile"
        assert snapshot["respondent"]["city"] == "Tokyo"

    category = authenticated.patch(
        f"/api/research/respondents/{created['id']}/category",
        headers=HEADERS,
        json={"final_category": "travel_ready"},
    )
    assert category.status_code == 200
    for item_id in interview_ids:
        history = authenticated.get(
            f"/api/research/interviews/{item_id}/versions", headers=HEADERS
        ).json()
        assert history["latest_version_number"] == 3
        latest = history["items"][0]
        snapshot = authenticated.get(
            f"/api/research/interviews/{item_id}/versions/{latest['id']}", headers=HEADERS
        ).json()["snapshot"]
        assert snapshot["respondent"]["final_category"] == "travel_ready"

    repeated_category = authenticated.patch(
        f"/api/research/respondents/{created['id']}/category",
        headers=HEADERS,
        json={"final_category": "travel_ready"},
    )
    assert repeated_category.status_code == 200
    assert authenticated.get(version_path, headers=HEADERS).json()["latest_version_number"] == 3


def test_respondent_list_search_filters_and_counts(authenticated):
    beginner = authenticated.post(
        "/api/research/respondents",
        headers=HEADERS,
        json={"screening_answers": {"name": "Лена Наблюдение", "independent_trips_24m": "0"}},
    ).json()
    experienced = authenticated.post(
        "/api/research/respondents",
        headers=HEADERS,
        json={"screening_answers": {"name": "Максим Поездки", "independent_trips_24m": "3_plus"}},
    ).json()
    interview_id = beginner["interviews"][0]["id"]
    authenticated.post(f"/api/research/interviews/{interview_id}/start", headers=HEADERS)
    authenticated.post(f"/api/research/interviews/{interview_id}/complete", headers=HEADERS)

    found = authenticated.get("/api/research/respondents?q=Наблюдение&page=1&page_size=1")
    assert found.status_code == 200
    assert found.json()["total"] == 1
    assert found.json()["items"][0]["id"] == beginner["id"]
    assert found.json()["items"][0]["status"] == "completed"

    filtered = authenticated.get(
        "/api/research/respondents?category=trips_experienced&status=draft"
    )
    assert filtered.json()["total"] == 1
    assert filtered.json()["items"][0]["id"] == experienced["id"]
    assert filtered.json()["total_count"] == 2
    assert filtered.json()["completed_count"] == 1
    assert filtered.json()["incomplete_count"] == 1


def test_concurrent_answer_cas_serializes_and_preserves_newer_text():
    token = uuid4().hex
    user = User(id=uuid4(), login=f"research-cas-{uuid4().hex}", password_hash="unused")
    with SessionLocal() as setup:
        setup.add(user)
        setup.flush()
        setup.add(
            AdminSession(
                token_digest=token_digest(token),
                user_id=user.id,
                expires_at=datetime.now(UTC) + timedelta(hours=1),
            )
        )
        setup.commit()

    def independent_db():
        with SessionLocal() as session:
            yield session

    application = create_app()
    application.dependency_overrides[get_db] = independent_db
    auth_headers = {**HEADERS, "Cookie": f"travel_session={token}"}
    respondent_id = None
    try:
        with TestClient(application) as setup_client:
            created_response = setup_client.post(
                "/api/research/respondents",
                headers=auth_headers,
                json={"screening_answers": {"name": "CAS concurrency"}},
            )
            assert created_response.status_code == 201, created_response.text
            created = created_response.json()
            respondent_id = created["id"]
            interview_id = created["interviews"][0]["id"]
            started = setup_client.post(
                f"/api/research/interviews/{interview_id}/start", headers=auth_headers
            )
            assert started.status_code == 200, started.text
            initial = setup_client.put(
                f"/api/research/interviews/{interview_id}/answers/G01",
                headers=auth_headers,
                json={"answer_text": "initial", "expected_revision": 0},
            )
            assert initial.status_code == 200, initial.text

        with TestClient(application) as first_client, TestClient(application) as second_client:
            clients = (first_client, second_client)

            def race_writes(expected_revision, texts):
                def write(index):
                    return clients[index].put(
                        f"/api/research/interviews/{interview_id}/answers/G01",
                        headers=auth_headers,
                        json={
                            "answer_text": texts[index],
                            "expected_revision": expected_revision,
                        },
                    )

                with ThreadPoolExecutor(max_workers=2) as executor:
                    return list(executor.map(write, (0, 1)))

            same_text_results = race_writes(1, ("shared retry", "shared retry"))
            assert [response.status_code for response in same_text_results] == [200, 200]
            assert {response.json()["revision"] for response in same_text_results} == {2}

            idempotent_retry = clients[0].put(
                f"/api/research/interviews/{interview_id}/answers/G01",
                headers=auth_headers,
                json={"answer_text": "shared retry", "expected_revision": 1},
            )
            assert idempotent_retry.status_code == 200
            assert idempotent_retry.json()["revision"] == 2

            competing_results = race_writes(2, ("writer one", "writer two"))
            successes = [response for response in competing_results if response.status_code == 200]
            conflicts = [response for response in competing_results if response.status_code == 409]
            assert len(successes) == len(conflicts) == 1, [
                (response.status_code, response.json()) for response in competing_results
            ]
            winner = successes[0].json()
            current = conflicts[0].json()["current"]
            assert current["revision"] == winner["revision"] == 3
            assert current["answer_text"] == winner["answer_text"]
    finally:
        with SessionLocal() as cleanup:
            if respondent_id:
                cleanup.execute(
                    delete(ResearchRespondent).where(
                        ResearchRespondent.id == respondent_id
                    )
                )
            cleanup.execute(delete(User).where(User.id == user.id))
            cleanup.commit()

def test_xlsx_export_contains_full_database_empty_answers_and_safe_long_text(authenticated):
    created = authenticated.post(
        "/api/research/respondents",
        headers=HEADERS,
        json={
            "screening_answers": {
                "name": "=1+1",
                "city": "Тверь",
                "departure_city": "Москва",
                "independent_trips_24m": "3_plus",
                "last_trip_organizer": "С друзьями",
            }
        },
    ).json()
    interview_id = created["interviews"][0]["id"]
    started = authenticated.post(
        f"/api/research/interviews/{interview_id}/start", headers=HEADERS
    ).json()
    assert started["scenario_code"] == "trips_experienced"
    long_text = ("Длинный ответ\n" * 3000) + ("x" * 32770)
    for code, text in (
        ("G01", '=HYPERLINK("https://example.invalid")\nстрока 2'),
        ("G02", long_text),
    ):
        saved = authenticated.put(
            f"/api/research/interviews/{interview_id}/answers/{code}",
            headers=HEADERS,
            json={"answer_text": text, "expected_revision": 0},
        )
        assert saved.status_code == 200

    response = authenticated.get("/api/research/export.xlsx")
    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content), data_only=False)
    assert workbook.sheetnames == [
        "Respondents",
        "Interviews",
        "Answers",
        "Questions",
        "Hypotheses",
    ]

    respondents = workbook["Respondents"]
    respondent_headers = [cell.value for cell in respondents[1]]
    assert respondents.cell(2, respondent_headers.index("Name") + 1).value == "=1+1"
    assert respondents.cell(2, respondent_headers.index("Name") + 1).data_type != "f"
    departure_city_header = "Screening: Откуда обычно вылетаете в путешествия?"
    last_trip_organizer_header = "Screening: Бронировали сами, с друзьями или через организатора?"
    assert departure_city_header in respondent_headers
    assert last_trip_organizer_header in respondent_headers
    assert "departure_city" not in respondent_headers
    assert "last_trip_organizer" not in respondent_headers
    assert (
        respondents.cell(2, respondent_headers.index(departure_city_header) + 1).value
        == "Москва"
    )
    assert (
        respondents.cell(2, respondent_headers.index(last_trip_organizer_header) + 1).value
        == "С друзьями"
    )

    answers = workbook["Answers"]
    answer_headers = [cell.value for cell in answers[1]]
    by_code = {
        row[answer_headers.index("Question Code")]: row
        for row in answers.iter_rows(min_row=2, values_only=True)
    }
    assert (
        by_code["G01"][answer_headers.index("Answer Text")]
        == '=HYPERLINK("https://example.invalid")\nстрока 2'
    )
    formula_cell = answers.cell(2, answer_headers.index("Answer Text") + 1)
    assert formula_cell.data_type != "f"
    assert "\n" in by_code["G01"][answer_headers.index("Answer Text")]
    assert by_code["G05"][answer_headers.index("Answer Text")] in {None, ""}
    exported_long = "".join(
        by_code["G02"][answer_headers.index("Answer Text") + offset] or ""
        for offset in range(max(1, (len(long_text) + 32766) // 32767))
    )
    assert exported_long == long_text
    assert by_code["G01"][answer_headers.index("Hypothesis IDs")]
    assert workbook["Hypotheses"].max_row == 29

    authenticated.put(
        f"/api/research/interviews/{interview_id}/answers/G01",
        headers=HEADERS,
        json={"answer_text": "Новая версия заметки", "expected_revision": 1},
    )
    refreshed = load_workbook(
        BytesIO(authenticated.get("/api/research/export.xlsx").content), data_only=False
    )
    refreshed_answers = refreshed["Answers"]
    refreshed_headers = [cell.value for cell in refreshed_answers[1]]
    refreshed_g01 = next(
        row
        for row in refreshed_answers.iter_rows(min_row=2, values_only=True)
        if row[refreshed_headers.index("Question Code")] == "G01"
    )
    assert refreshed_g01[refreshed_headers.index("Answer Text")] == "Новая версия заметки"
